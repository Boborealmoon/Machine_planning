"""APS process-sheet Excel matcher: headers, PS duplication, inventory pick, statuses."""
from __future__ import annotations

import io
import json
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from app import app
from planning.aps_ps_match_service import (
    FIELD_LABELS,
    compact_lot_key,
    field_for_header,
    index_lot_rows,
    index_outstanding_aps,
    is_outstanding_cache_row,
    lot_lookup_keys,
    lots_for_in_house_ref,
    match_source_rows,
    normalize_header,
    parse_inventory_dimensions,
    parse_workbook,
    pick_inventory,
    process_upload,
    score_inventory_candidate,
    summarize_rows,
)


def _xlsx_bytes(headers, rows, *, two_line=False, extra_header=None):
    book = Workbook()
    ws = book.active
    ws.title = "APS"
    if two_line:
        top, bottom = extra_header
        ws.append(top)
        ws.append(bottom)
    else:
        ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def _source(**overrides):
    row = {
        "part_no": "88D012",
        "description": "Bush",
        "material_type": "Bronze",
        "in_mm": 38.1,
        "bush_flange_od_max": 50,
        "bush_id_min": 30,
        "flange_bush_length": 40,
        "length_inch_per_piece": 2.5,
        "total_length_inch": 10,
        "our_matl_type": "CuZn19Al6",
        "matl_od": 38.1,
        "matl_id": "",
        "in_house_ref": "AM/0454/21",
    }
    row.update(overrides)
    return row


class HeaderMappingTests(unittest.TestCase):
    def test_normalize_header_collapses_newlines_and_apostrophes(self):
        self.assertEqual(normalize_header("In-house M/AM Ref"), "in_house_m_am_ref")
        self.assertEqual(normalize_header("Our MAT'L Type"), "our_mat_l_type")
        self.assertEqual(
            normalize_header("Length (inch\nPer piece"),
            "length_inch_per_piece",
        )
        self.assertEqual(
            normalize_header("Total Length (Inch)\nRequired per line item"),
            "total_length_inch_required_per_line_item",
        )

    def test_field_aliases(self):
        self.assertEqual(field_for_header("Part Number"), "part_no")
        self.assertEqual(field_for_header("In-house M/AM Ref"), "in_house_ref")
        self.assertEqual(field_for_header("Our MAT'L Type"), "our_matl_type")
        self.assertEqual(field_for_header("MAT'L OD"), "matl_od")
        self.assertEqual(field_for_header("Length (inch) Per piece"), "length_inch_per_piece")
        self.assertEqual(
            field_for_header("Total Length (Inch) Required per line item"),
            "total_length_inch",
        )

    def test_parse_wrapped_single_header_row(self):
        payload = _xlsx_bytes(
            [
                "Part Number",
                "Description",
                "Our MAT'L Type",
                "MAT'L OD",
                "MAT'L ID",
                "In-house M/AM Ref",
            ],
            [["88D012", "Bush", "CuZn19Al6", 38.1, "", "AM/0454/21"]],
        )
        rows = parse_workbook(payload, "aps.xlsx")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["part_no"], "88D012")
        self.assertEqual(rows[0]["in_house_ref"], "AM/0454/21")
        self.assertEqual(rows[0]["our_matl_type"], "CuZn19Al6")
        self.assertEqual(rows[0]["matl_od"], 38.1)

    def test_parse_two_line_headers(self):
        payload = _xlsx_bytes(
            [],
            [["88D012", "Bush", 2.5, 10, "AM/0454/21"]],
            two_line=True,
            extra_header=(
                ["Part Number", "Description", "Length (inch", "Total Length (Inch)", "In-house M/AM Ref"],
                ["", "", "Per piece", "Required per line item", ""],
            ),
        )
        rows = parse_workbook(payload, "aps.xlsx")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["part_no"], "88D012")
        self.assertEqual(rows[0]["length_inch_per_piece"], 2.5)
        self.assertEqual(rows[0]["total_length_inch"], 10)
        self.assertEqual(rows[0]["in_house_ref"], "AM/0454/21")


class OutstandingApsTests(unittest.TestCase):
    def test_historical_and_shipped_are_not_active(self):
        self.assertFalse(is_outstanding_cache_row({"so_det_qty": None, "qty_shipped": 0}))
        self.assertFalse(is_outstanding_cache_row({"so_det_qty": "", "qty_shipped": 0}))
        self.assertTrue(is_outstanding_cache_row({"so_det_qty": 10, "qty_shipped": 9.9}))
        self.assertFalse(is_outstanding_cache_row({"so_det_qty": 10, "qty_shipped": 10}))

    def test_index_keeps_active_aps_and_drops_historical(self):
        by_part = index_outstanding_aps(
            [
                {"ps_id": "APS26-0151", "part_no": "88D012", "so_det_qty": 10, "qty_shipped": 0},
                {"ps_id": "APS26-0151", "part_no": "88D012", "so_det_qty": 10, "qty_shipped": 0},
                {"ps_id": "NPS26-0001", "part_no": "88D012", "so_det_qty": 10, "qty_shipped": 0},
                {"ps_id": "APS26-0203", "part_no": "88D012", "so_det_qty": 4, "qty_shipped": 0},
                {"ps_id": "APS26-0999", "part_no": "88D012", "so_det_qty": 4, "qty_shipped": 4},
                {"ps_id": "APS26-0100", "part_no": "88D012", "so_det_qty": None, "qty_shipped": 0},
                {"ps_id": "APS26-0300", "part_no": "OTHER", "so_det_qty": 1, "qty_shipped": 0},
            ]
        )
        self.assertEqual(by_part["88D012"], ["APS26-0151", "APS26-0203"])
        self.assertEqual(by_part["OTHER"], ["APS26-0300"])
        self.assertNotIn("APS26-0100", by_part.get("88D012", []))
        self.assertNotIn("APS26-0999", by_part.get("88D012", []))


class InventoryPickTests(unittest.TestCase):
    def test_dimension_suffix_parse(self):
        self.assertEqual(parse_inventory_dimensions("CuZn19Al6*3_D38.1"), (38.1, None))
        self.assertEqual(parse_inventory_dimensions("NITRONIC 50(HS)*3_D50.8_39.1"), (50.8, 39.1))

    def test_lot_keys_include_batch_and_compact_am_ref(self):
        keys = lot_lookup_keys("AM/0454/21")
        self.assertIn("AM/0454/21", keys)
        self.assertIn("AM045421", keys)
        self.assertIn("AM45421", keys)
        self.assertEqual(compact_lot_key("AM/0454/21"), compact_lot_key("AM/454/21"))
        self.assertIn("454", lot_lookup_keys(454))
        self.assertIn("454", lot_lookup_keys(454.0))

    def test_reverse_search_by_batch_no(self):
        indexed = index_lot_rows(
            [
                {
                    "inventory_code": "CuZn19Al6*3_D38.1",
                    "reference_no": "AM/0454/21",
                    "lot_no": "454",
                    "remaining_qty": 12,
                    "inventory_class_code": "RAW MATERIAL",
                }
            ]
        )
        by_batch = lots_for_in_house_ref("454", indexed)
        self.assertEqual(by_batch[0]["inventory_code"], "CuZn19Al6*3_D38.1")
        by_ref = lots_for_in_house_ref("AM/0454/21", indexed)
        self.assertEqual(by_ref[0]["inventory_code"], "CuZn19Al6*3_D38.1")
        by_compact = lots_for_in_house_ref("AM/454/21", indexed)
        self.assertEqual(by_compact[0]["inventory_code"], "CuZn19Al6*3_D38.1")

    def test_match_uses_batch_no_when_ref_key_missing(self):
        matched = match_source_rows(
            [_source(in_house_ref="454")],
            {"88D012": ["APS26-0151"]},
            {
                "454": [
                    {
                        "inventory_code": "CuZn19Al6*3_D38.1",
                        "batch_no": "454",
                        "inventory_class_code": "RAW MATERIAL",
                        "remaining_qty": 12,
                    }
                ]
            },
        )
        self.assertEqual(matched[0]["inventory_code"], "CuZn19Al6*3_D38.1")
        self.assertEqual(matched[0]["match_status"], "ok")

    def test_prefers_raw_material_and_od_match(self):
        source = _source()
        fg = {
            "inventory_code": "88D012",
            "inventory_class_code": "FG MFG COMMERCIAL",
            "remaining_qty": 20,
        }
        raw = {
            "inventory_code": "CuZn19Al6*3_D38.1",
            "inventory_class_code": "RAW MATERIAL",
            "remaining_qty": 12,
        }
        self.assertGreater(
            score_inventory_candidate(raw, source),
            score_inventory_candidate(fg, source),
        )
        picked, ambiguous = pick_inventory([fg, raw], source)
        self.assertEqual(picked["inventory_code"], "CuZn19Al6*3_D38.1")
        self.assertFalse(ambiguous)

    def test_ambiguous_when_two_raw_codes_tie(self):
        source = _source(our_matl_type="", matl_od="", matl_id="")
        first = {
            "inventory_code": "CuZn19Al6*3_D38.1",
            "inventory_class_code": "RAW MATERIAL",
            "remaining_qty": 5,
        }
        second = {
            "inventory_code": "CuZn19Al6*3_D50.8",
            "inventory_class_code": "RAW MATERIAL",
            "remaining_qty": 5,
        }
        picked, ambiguous = pick_inventory([first, second], source)
        self.assertTrue(ambiguous)
        self.assertIsNotNone(picked)


class MatchRowTests(unittest.TestCase):
    def test_duplicates_excel_row_for_multiple_ps(self):
        matched = match_source_rows(
            [_source()],
            {"88D012": ["APS26-0151", "APS26-0203"]},
            {
                "AM/0454/21": [
                    {
                        "inventory_code": "CuZn19Al6*3_D38.1",
                        "inventory_class_code": "RAW MATERIAL",
                        "remaining_qty": 12,
                    }
                ]
            },
        )
        self.assertEqual(len(matched), 2)
        self.assertEqual([row["matched_ps_no"] for row in matched], ["APS26-0151", "APS26-0203"])
        self.assertTrue(all(row["inventory_code"] == "CuZn19Al6*3_D38.1" for row in matched))
        self.assertTrue(all(row["match_status"] == "ok" for row in matched))

    def test_unmatched_statuses(self):
        both = match_source_rows(
            [_source(part_no="MISSING", in_house_ref="")],
            {},
            {},
        )
        self.assertEqual(both[0]["match_status"], "ps_and_inventory_unmatched")
        ps_only = match_source_rows(
            [_source(in_house_ref="AM/NONE/00")],
            {"88D012": ["APS26-0151"]},
            {},
        )
        self.assertEqual(ps_only[0]["match_status"], "unmatched_inventory")
        self.assertEqual(ps_only[0]["matched_ps_no"], "APS26-0151")
        inv_only = match_source_rows(
            [_source(part_no="GONE")],
            {},
            {
                "AM/0454/21": [
                    {
                        "inventory_code": "CuZn19Al6*3_D38.1",
                        "inventory_class_code": "RAW MATERIAL",
                        "remaining_qty": 1,
                    }
                ]
            },
        )
        self.assertEqual(inv_only[0]["match_status"], "unmatched_ps")
        self.assertEqual(inv_only[0]["inventory_code"], "CuZn19Al6*3_D38.1")

    def test_inventory_ambiguous_status(self):
        matched = match_source_rows(
            [_source(our_matl_type="", matl_od="", matl_id="")],
            {"88D012": ["APS26-0151"]},
            {
                "AM/0454/21": [
                    {
                        "inventory_code": "CuZn19Al6*3_D38.1",
                        "inventory_class_code": "RAW MATERIAL",
                        "remaining_qty": 5,
                    },
                    {
                        "inventory_code": "CuZn19Al6*3_D50.8",
                        "inventory_class_code": "RAW MATERIAL",
                        "remaining_qty": 5,
                    },
                ]
            },
        )
        self.assertEqual(matched[0]["match_status"], "inventory_ambiguous")
        self.assertTrue(matched[0]["inventory_code"])

    def test_summary_counts(self):
        summary = summarize_rows(
            [
                {"matched_ps_no": "APS26-0151", "inventory_code": "A", "match_status": "ok"},
                {"matched_ps_no": "", "inventory_code": "", "match_status": "ps_and_inventory_unmatched"},
            ]
        )
        self.assertEqual(summary["output_rows"], 2)
        self.assertEqual(summary["matched_ps"], 1)
        self.assertEqual(summary["unmatched_ps"], 1)
        self.assertEqual(summary["unmatched_inventory"], 1)


class ProcessUploadAndRouteTests(unittest.TestCase):
    def test_process_upload_writes_export_columns(self):
        payload = _xlsx_bytes(
            ["Part Number", "In-house M/AM Ref"],
            [["88D012", "AM/0454/21"]],
        )
        with patch("planning.aps_ps_match_service.fetch_outstanding_aps", return_value={"88D012": ["APS26-0151"]}), patch(
            "planning.aps_ps_match_service.fetch_lots_by_reference_nos",
            return_value={
                "AM/0454/21": [
                    {
                        "inventory_code": "CuZn19Al6*3_D38.1",
                        "inventory_class_code": "RAW MATERIAL",
                        "remaining_qty": 12,
                    }
                ]
            },
        ), patch("planning.aps_ps_match_service.fetch_inventory_classes", return_value={}):
            workbook_bytes, summary = process_upload(payload, "aps.xlsx")
        self.assertEqual(summary["matched_ps"], 1)
        self.assertEqual(summary["input_rows"], 1)
        book = load_workbook(io.BytesIO(workbook_bytes))
        ws = book.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn(FIELD_LABELS["matched_ps_no"], headers)
        self.assertIn(FIELD_LABELS["inventory_code"], headers)
        values = [cell.value for cell in ws[2]]
        self.assertIn("APS26-0151", values)
        self.assertIn("CuZn19Al6*3_D38.1", values)

    def test_route_returns_xlsx(self):
        payload = _xlsx_bytes(["Part Number", "In-house M/AM Ref"], [["88D012", "AM/0454/21"]])
        fake_book = _xlsx_bytes(["Matched PS No"], [["APS26-0151"]])
        summary = {
            "ok": True,
            "input_rows": 1,
            "output_rows": 1,
            "matched_ps": 1,
            "unmatched_ps": 0,
            "unmatched_inventory": 0,
            "download_name": "aps-ps-match-20260910.xlsx",
        }
        with patch.dict("os.environ", {"PLANNER_PASSCODE": "", "ADMIN_PASSCODE": ""}), app.test_client() as client, patch(
            "planning.aps_ps_match_route.process_upload",
            return_value=(fake_book, summary),
        ):
            response = client.post(
                "/api/archive/aps-ps-match",
                data={"file": (io.BytesIO(payload), "aps.xlsx")},
                content_type="multipart/form-data",
            )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response.content_type)
        parsed = json.loads(response.headers["X-Aps-Match-Summary"])
        self.assertEqual(parsed["matched_ps"], 1)

    def test_page_renders(self):
        with patch.dict("os.environ", {"PLANNER_PASSCODE": "", "ADMIN_PASSCODE": ""}), app.test_client() as client:
            response = client.get("/archive/aps-ps-match")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"APS PS material match", response.data)

    def test_admin_hub_links_to_page(self):
        with patch.dict("os.environ", {"PLANNER_PASSCODE": "", "ADMIN_PASSCODE": ""}), app.test_client() as client:
            response = client.get("/admin")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'href="/archive/aps-ps-match"', response.data)
