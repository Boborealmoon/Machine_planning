"""APS process-sheet Excel matcher: parse order Excel, match outstanding APS, look up lots."""
from __future__ import annotations

import io
import logging
import re
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from db import planner_db_connect_error

from .utils import SHIPPED_QTY_TOLERANCE, compact_text, parse_number, planner_today

logger = logging.getLogger(__name__)

MAX_UPLOAD_BYTES = 12 * 1024 * 1024
MAX_SOURCE_ROWS = 5000
HEADER_SCAN_ROWS = 20
DIM_TOLERANCE = 0.051

SOURCE_FIELDS = (
    "part_no",
    "description",
    "material_type",
    "in_mm",
    "bush_flange_od_max",
    "bush_id_min",
    "flange_bush_length",
    "length_inch_per_piece",
    "total_length_inch",
    "our_matl_type",
    "matl_od",
    "matl_id",
    "in_house_ref",
)

EXPORT_FIELDS = (
    *SOURCE_FIELDS,
    "matched_ps_no",
    "inventory_code",
    "lot_remaining_qty",
    "match_status",
)

FIELD_LABELS = {
    "part_no": "Part Number",
    "description": "Description",
    "material_type": "Material Type",
    "in_mm": "in MM",
    "bush_flange_od_max": "Bush/Flange OD Max",
    "bush_id_min": "Bush ID Min",
    "flange_bush_length": "Flange + Bush Length",
    "length_inch_per_piece": "Length (inch) Per piece",
    "total_length_inch": "Total Length (Inch) Required per line item",
    "our_matl_type": "Our MAT'L Type",
    "matl_od": "MAT'L OD",
    "matl_id": "MAT'L ID",
    "in_house_ref": "In-house M/AM Ref",
    "matched_ps_no": "Matched PS No",
    "inventory_code": "Inventory Code",
    "lot_remaining_qty": "Lot remaining qty",
    "match_status": "Match status",
}

_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "part_no": (
        "part_no",
        "part_number",
        "partno",
        "partnumber",
        "item_code",
        "item_no",
        "drawing_no",
    ),
    "description": ("description", "desc", "part_description", "part_desc"),
    "material_type": ("material_type", "material", "mat_type"),
    "in_mm": ("in_mm", "mm", "diameter_mm"),
    "bush_flange_od_max": (
        "bush_flange_od_max",
        "bush_flange_od",
        "bush_od_max",
        "flange_od_max",
        "bush_flange_od_max_mm",
    ),
    "bush_id_min": ("bush_id_min", "bush_id", "bush_id_min_mm"),
    "flange_bush_length": (
        "flange_bush_length",
        "flange_bush_length_mm",
        "flange_plus_bush_length",
    ),
    "length_inch_per_piece": (
        "length_inch_per_piece",
        "length_inch",
        "length_per_piece",
        "length_inch_per_piece",
        "length_inch_per_piece",
    ),
    "total_length_inch": (
        "total_length_inch",
        "total_length_inch_required_per_line_item",
        "total_length",
        "total_length_required_per_line_item",
        "total_length_inch_required_per_line_item",
    ),
    "our_matl_type": (
        "our_mat_l_type",
        "our_matl_type",
        "our_material_type",
        "our_mat_l",
        "our_matl",
    ),
    "matl_od": ("mat_l_od", "matl_od", "material_od", "mat_od"),
    "matl_id": ("mat_l_id", "matl_id", "material_id", "material_id_od"),
    "in_house_ref": (
        "in_house_m_am_ref",
        "in_house_mam_ref",
        "in_house_ref",
        "inhouse_m_am_ref",
        "in_house_m_am",
        "lot_no",
        "lot_ref",
        "lot_reference_no",
        "batch_no",
        "am_ref",
    ),
}

_DIM_RE = re.compile(
    r"_D(?P<od>\d+(?:\.\d+)?)(?:_(?P<id>\d+(?:\.\d+)?))?\s*$",
    re.I,
)
_NUMBER_RE = re.compile(r"-?\d+(?:\.\d+)?")
_HEADER_HINTS = (
    "part_number",
    "part_no",
    "in_house",
    "mat_l",
    "bush_flange",
    "our_mat",
)

# Live Sales Orders "Active" APS: outstanding SO lines only (so_order_ost_det).
# Historical / fully shipped jobs drop off ost_det and must not be matched.
_ACTIVE_APS_SQL = f"""
SELECT
    UPPER(BTRIM(split_part(
        COALESCE(NULLIF(BTRIM(ps.process_sheet_no), ''), pp.pp_voucher_no),
        '::',
        1
    ))) AS ps_id,
    MAX(COALESCE(NULLIF(BTRIM(ps.inventory_code), ''), pp.inventory_code)) AS part_no,
    MAX(det.qty) AS so_det_qty,
    MAX(COALESCE(sq.qty_shipped, 0)) AS qty_shipped
FROM public.mfg_pp_vch pp
LEFT JOIN public.mfg_process_sheet_info_v1_view ps
       ON ps.pp_voucher_no = pp.pp_voucher_no
INNER JOIN public.so_order_ost_det det
        ON det.sales_order_no = pp.source_voucher_no
       AND regexp_replace(det.line_item_no::TEXT, '\\.0+$', '')
           = regexp_replace(pp.source_line_item_no::TEXT, '\\.0+$', '')
LEFT JOIN public.sum_qty_shipped_by_sales_order sq
       ON sq.sales_order_no = pp.source_voucher_no
      AND regexp_replace(sq.line_item_no::TEXT, '\\.0+$', '')
          = regexp_replace(pp.source_line_item_no::TEXT, '\\.0+$', '')
WHERE pp.source_voucher_no IS NOT NULL
  AND starts_with(
        UPPER(BTRIM(COALESCE(NULLIF(BTRIM(ps.process_sheet_no), ''), pp.pp_voucher_no))),
        'APS'
      )
  AND BTRIM(COALESCE(NULLIF(BTRIM(ps.inventory_code), ''), pp.inventory_code, '')) <> ''
  AND COALESCE(sq.qty_shipped, 0) < det.qty - {SHIPPED_QTY_TOLERANCE}
GROUP BY 1
ORDER BY 1
"""

_LOTS_BY_BATCH_SQL = """
SELECT
    o.inventory_code,
    o.reference_no,
    o.lot_no,
    SUM(COALESCE(o.remaining_qty, 0)) AS remaining_qty,
    SUM(COALESCE(o.available_qty, 0)) AS available_qty,
    SUM(COALESCE(o.allocation_qty, 0)) AS allocation_qty
FROM public.ic_inventory_ost_lot o
WHERE o.inventory_code IS NOT NULL
  AND BTRIM(o.inventory_code) <> ''
  AND (
    UPPER(REPLACE(BTRIM(COALESCE(o.reference_no, '')), ' ', '')) = ANY(%s)
    OR UPPER(REPLACE(BTRIM(COALESCE(o.lot_no::TEXT, '')), ' ', '')) = ANY(%s)
    OR UPPER(REGEXP_REPLACE(COALESCE(o.reference_no, ''), '[^A-Za-z0-9]', '', 'g')) = ANY(%s)
    OR UPPER(REGEXP_REPLACE(COALESCE(o.lot_no::TEXT, ''), '[^A-Za-z0-9]', '', 'g')) = ANY(%s)
    OR EXISTS (
        SELECT 1
        FROM unnest(%s::text[]) AS needle
        WHERE length(needle) >= 4
          AND (
            UPPER(REPLACE(BTRIM(COALESCE(o.reference_no, '')), ' ', '')) LIKE '%%' || needle || '%%'
            OR UPPER(REPLACE(BTRIM(COALESCE(o.lot_no::TEXT, '')), ' ', '')) LIKE '%%' || needle || '%%'
            OR UPPER(REGEXP_REPLACE(COALESCE(o.reference_no, ''), '[^A-Za-z0-9]', '', 'g'))
                LIKE '%%' || needle || '%%'
            OR UPPER(REGEXP_REPLACE(COALESCE(o.lot_no::TEXT, ''), '[^A-Za-z0-9]', '', 'g'))
                LIKE '%%' || needle || '%%'
          )
    )
  )
GROUP BY o.inventory_code, o.reference_no, o.lot_no
ORDER BY o.inventory_code, o.reference_no, o.lot_no
"""

_CLASS_BY_CODE_SQL = """
SELECT inventory_code, inventory_class_code
FROM public.ic_inventory_enquiry_summary_view
WHERE inventory_code IS NOT NULL
  AND BTRIM(inventory_code) <> ''
  AND UPPER(BTRIM(inventory_code)) = ANY(%s)
"""


def json_error(exc: Exception, *, fallback_status: int = 500):
    friendly = planner_db_connect_error(exc)
    if friendly:
        return {"error": friendly}, 503
    text = str(exc) or exc.__class__.__name__
    lower = text.lower()
    if (
        "timed out" in lower
        or "timeout" in lower
        or "querycanceled" in lower
        or "statement timeout" in lower
        or "read timed out" in lower
    ):
        return {"error": "APS match timed out while reading process sheets or inventory."}, 504
    return {"error": text}, fallback_status


def normalize_header(value: Any) -> str:
    text = compact_text(value).lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def normalize_part_no(value: Any) -> str:
    text = compact_text(value).upper()
    text = re.sub(r"[\s_]+", "-", text)
    text = re.sub(r"-{2,}", "-", text)
    return text.strip("-")


_REF_SPLIT_RE = re.compile(r"[,;|\n]+")
_LEADING_ZERO_RE = re.compile(r"(?<=[A-Z])0+(?=\d)")


def normalize_lot_ref(value: Any) -> str:
    if isinstance(value, float) and value == int(value):
        value = int(value)
    text = compact_text(value).upper()
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    return re.sub(r"\s+", "", text)


def compact_lot_key(value: Any) -> str:
    alnum = re.sub(r"[^A-Z0-9]+", "", normalize_lot_ref(value))
    return _LEADING_ZERO_RE.sub("", alnum)


def lot_lookup_keys(value: Any) -> list[str]:
    """Keys for reverse lot lookup: In-house ref, batch/lot no, and compact forms."""
    keys: list[str] = []
    seen: set[str] = set()

    def add(item: Any) -> None:
        text = normalize_lot_ref(item)
        if not text or text in seen:
            return
        seen.add(text)
        keys.append(text)
        alnum = re.sub(r"[^A-Z0-9]+", "", text)
        if alnum and alnum not in seen:
            seen.add(alnum)
            keys.append(alnum)
        stripped = compact_lot_key(text)
        if stripped and stripped not in seen:
            seen.add(stripped)
            keys.append(stripped)

    raw = compact_text(value)
    if isinstance(value, float) and value == int(value):
        add(int(value))
    parts = [part.strip() for part in _REF_SPLIT_RE.split(raw)] if raw else []
    if not parts:
        add(value)
    for part in parts:
        add(part)
    return keys


def _alnum_key(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "", compact_text(value).upper())


def parse_dimension_number(value: Any) -> float | None:
    text = compact_text(value).replace(",", "")
    match = _NUMBER_RE.search(text)
    if not match:
        return None
    try:
        return float(match.group(0))
    except (TypeError, ValueError):
        return None


def parse_inventory_dimensions(inventory_code: str) -> tuple[float | None, float | None]:
    match = _DIM_RE.search(compact_text(inventory_code))
    if not match:
        return None, None
    od = float(match.group("od")) if match.group("od") else None
    id_val = float(match.group("id")) if match.group("id") else None
    return od, id_val


def _nums_close(left: float | None, right: float | None, tol: float = DIM_TOLERANCE) -> bool:
    if left is None or right is None:
        return False
    return abs(left - right) <= tol


def field_for_header(header: str, *, used: set[str] | None = None) -> str:
    norm = normalize_header(header)
    if not norm:
        return ""
    taken = used or set()
    for field, aliases in _HEADER_ALIASES.items():
        if field in taken:
            continue
        if norm == field or norm in aliases:
            return field
        if field in {"length_inch_per_piece", "total_length_inch", "in_house_ref"}:
            if any(alias in norm and alias for alias in aliases if len(alias) >= 8):
                if field == "length_inch_per_piece" and "total" in norm:
                    continue
                return field
    return ""


def _serialize_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, float) and value != value:
        return ""
    return value


def _row_has_values(row: list[Any]) -> bool:
    return any(compact_text(cell) for cell in (row or []))


def _merge_header_cells(first: list[Any], second: list[Any]) -> list[str]:
    width = max(len(first or []), len(second or []))
    merged: list[str] = []
    for idx in range(width):
        top = compact_text(first[idx] if idx < len(first) else "")
        bottom = compact_text(second[idx] if idx < len(second) else "")
        if top and bottom:
            merged.append(f"{top} {bottom}")
        else:
            merged.append(top or bottom)
    return merged


def _mapped_fields(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    used: set[str] = set()
    for idx, header in enumerate(headers):
        field = field_for_header(header, used=used)
        if not field or field in mapping:
            continue
        mapping[field] = idx
        used.add(field)
    return mapping


def _header_score(headers: list[str]) -> int:
    mapping = _mapped_fields(headers)
    score = len(mapping)
    if "part_no" in mapping:
        score += 8
    if "in_house_ref" in mapping:
        score += 4
    return score


def _looks_like_header_row(row: list[Any]) -> bool:
    joined = " ".join(normalize_header(cell) for cell in (row or []) if compact_text(cell))
    if not joined:
        return False
    return any(hint in joined for hint in _HEADER_HINTS)


def _find_header(matrix: list[list[Any]]) -> tuple[int, int, list[str], dict[str, int]]:
    best_idx = 0
    best_span = 1
    best_headers: list[str] = []
    best_map: dict[str, int] = {}
    best_score = -1
    limit = min(len(matrix), HEADER_SCAN_ROWS)
    for idx in range(limit):
        row = matrix[idx] or []
        if not _row_has_values(row):
            continue
        single = [compact_text(cell) for cell in row]
        candidates = [(single, 1)]
        if idx + 1 < len(matrix) and _looks_like_header_row(row):
            nxt = matrix[idx + 1] or []
            merged = _merge_header_cells(row, nxt)
            if len(_mapped_fields(merged)) > len(_mapped_fields(single)):
                candidates.append((merged, 2))
        for headers, span in candidates:
            mapping = _mapped_fields(headers)
            score = _header_score(headers)
            if "part_no" not in mapping:
                continue
            if score > best_score:
                best_idx = idx
                best_span = span
                best_headers = headers
                best_map = mapping
                best_score = score
    if "part_no" not in best_map:
        raise ValueError(
            "Could not find a Part Number column. Use a sheet with Part Number and In-house M/AM Ref headers."
        )
    return best_idx, best_span, best_headers, best_map


def _sheet_matrix_from_xlsx(payload: bytes) -> list[tuple[str, list[list[Any]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(payload), data_only=True)
    sheets: list[tuple[str, list[list[Any]]]] = []
    try:
        for ws in workbook.worksheets:
            try:
                ws.reset_dimensions()
            except Exception:
                pass
            matrix: list[list[Any]] = []
            for row in ws.iter_rows(values_only=True):
                matrix.append([_serialize_cell(cell) for cell in row])
            sheets.append((compact_text(ws.title) or "Sheet1", matrix))
    finally:
        workbook.close()
    return sheets


def _sheet_matrix_from_xls(payload: bytes) -> list[tuple[str, list[list[Any]]]]:
    try:
        import xlrd
    except Exception as exc:
        raise RuntimeError("xlrd is required to read .xls files.") from exc
    book = xlrd.open_workbook(file_contents=payload)
    sheets: list[tuple[str, list[list[Any]]]] = []
    for sheet in book.sheets():
        matrix: list[list[Any]] = []
        for row_idx in range(sheet.nrows):
            matrix.append(
                [_serialize_cell(sheet.cell_value(row_idx, col_idx)) for col_idx in range(sheet.ncols)]
            )
        sheets.append((compact_text(sheet.name) or "Sheet1", matrix))
    return sheets


def _source_rows_from_matrix(matrix: list[list[Any]]) -> list[dict[str, Any]]:
    header_idx, header_span, _headers, mapping = _find_header(matrix)
    data_start = header_idx + header_span
    records: list[dict[str, Any]] = []
    for row in matrix[data_start:]:
        if not _row_has_values(row or []):
            continue
        record: dict[str, Any] = {field: "" for field in SOURCE_FIELDS}
        for field, idx in mapping.items():
            record[field] = row[idx] if idx < len(row) else ""
        if not compact_text(record.get("part_no")):
            continue
        records.append(record)
        if len(records) >= MAX_SOURCE_ROWS:
            break
    return records


def parse_workbook(payload: bytes, filename: str = "") -> list[dict[str, Any]]:
    if not payload:
        raise ValueError("The Excel file is empty")
    if len(payload) > MAX_UPLOAD_BYTES:
        raise ValueError("Excel file is larger than 12 MB")
    name = compact_text(filename).lower()
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        sheets = _sheet_matrix_from_xls(payload)
    else:
        sheets = _sheet_matrix_from_xlsx(payload)
    best: list[dict[str, Any]] = []
    last_error: Exception | None = None
    for _title, matrix in sheets:
        try:
            items = _source_rows_from_matrix(matrix)
        except ValueError as exc:
            last_error = exc
            continue
        if len(items) > len(best):
            best = items
    if not best:
        if last_error:
            raise last_error
        raise ValueError("No part-number rows were found in that workbook.")
    return best


def is_active_aps_row(row: dict[str, Any]) -> bool:
    """Active APS = still on an outstanding SO line and not fully shipped.

    Missing SO qty means historical / closed (not on so_order_ost_det).
    """
    so_qty = row.get("so_det_qty")
    if so_qty in (None, ""):
        return False
    try:
        required = float(so_qty)
    except (TypeError, ValueError):
        return False
    shipped = parse_number(row.get("qty_shipped"), 0)
    return shipped < required - SHIPPED_QTY_TOLERANCE


is_outstanding_cache_row = is_active_aps_row


def index_outstanding_aps(cache_rows: list[dict[str, Any]]) -> dict[str, list[str]]:
    by_part: dict[str, list[str]] = {}
    seen: set[tuple[str, str]] = set()
    for row in cache_rows:
        ps_id = compact_text(row.get("ps_id")).upper()
        if "::" in ps_id:
            ps_id = ps_id.split("::", 1)[0]
        if not ps_id.startswith("APS"):
            continue
        if not is_active_aps_row(row):
            continue
        part = normalize_part_no(row.get("part_no"))
        if not part:
            continue
        key = (ps_id, part)
        if key in seen:
            continue
        seen.add(key)
        by_part.setdefault(part, []).append(ps_id)
    for part, ps_ids in by_part.items():
        by_part[part] = sorted(dict.fromkeys(ps_ids))
    return by_part


def fetch_outstanding_aps() -> dict[str, list[str]]:
    from .staged_erp import live_query

    return index_outstanding_aps(live_query(_ACTIVE_APS_SQL))


def index_lot_rows(lot_rows: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    """Index lots by reference_no and batch/lot_no so Excel can reverse-search either."""
    lots_by_key: dict[str, list[dict[str, Any]]] = {}
    seen: set[tuple[str, str, str, str]] = set()
    for row in lot_rows:
        code = compact_text(row.get("inventory_code"))
        if not code:
            continue
        item = {
            "inventory_code": code,
            "reference_no": compact_text(row.get("reference_no")),
            "batch_no": compact_text(row.get("batch_no") or row.get("lot_no")),
            "lot_no": compact_text(row.get("lot_no") or row.get("batch_no")),
            "remaining_qty": parse_number(row.get("remaining_qty"), 0),
            "available_qty": parse_number(row.get("available_qty"), 0),
            "allocation_qty": parse_number(row.get("allocation_qty"), 0),
            "inventory_class_code": compact_text(row.get("inventory_class_code")),
        }
        for key in (
            *lot_lookup_keys(item.get("reference_no")),
            *lot_lookup_keys(item.get("batch_no")),
        ):
            sig = (key, code, item["reference_no"], item["batch_no"])
            if sig in seen:
                continue
            seen.add(sig)
            lots_by_key.setdefault(key, []).append(item)
    return lots_by_key


def lots_for_in_house_ref(
    in_house_ref: Any,
    lots_by_key: dict[str, list[dict[str, Any]]],
) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()

    def add(item: dict[str, Any]) -> None:
        sig = (
            compact_text(item.get("inventory_code")).upper(),
            compact_text(item.get("reference_no")),
            compact_text(item.get("batch_no") or item.get("lot_no")),
        )
        if not sig[0] or sig in seen:
            return
        seen.add(sig)
        candidates.append(item)

    needles = lot_lookup_keys(in_house_ref)
    for key in needles:
        for item in lots_by_key.get(key) or []:
            add(item)
    if candidates or not needles:
        return candidates
    for item in _unique_lot_items(lots_by_key):
        hay = {
            *lot_lookup_keys(item.get("reference_no")),
            *lot_lookup_keys(item.get("batch_no") or item.get("lot_no")),
        }
        for needle in needles:
            if needle in hay:
                add(item)
                break
            if len(needle) >= 4 and any(needle in key or key in needle for key in hay if key):
                add(item)
                break
    return candidates


def _unique_lot_items(lots_by_key: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    unique: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    for items in lots_by_key.values():
        for item in items:
            sig = (
                compact_text(item.get("inventory_code")).upper(),
                compact_text(item.get("reference_no")),
                compact_text(item.get("batch_no") or item.get("lot_no")),
            )
            if not sig[0] or sig in seen:
                continue
            seen.add(sig)
            unique.append(item)
    return unique


def fetch_lots_by_reference_nos(refs: list[str]) -> dict[str, list[dict[str, Any]]]:
    keys = sorted({key for ref in refs for key in lot_lookup_keys(ref)})
    if not keys:
        return {}
    needles = [key for key in keys if len(key) >= 4]
    from .staged_erp import live_query

    rows_out = live_query(_LOTS_BY_BATCH_SQL, (keys, keys, keys, keys, needles))
    return index_lot_rows(rows_out)


def fetch_inventory_classes(codes: list[str]) -> dict[str, str]:
    normalized = sorted({compact_text(code).upper() for code in codes if compact_text(code)})
    if not normalized:
        return {}
    from .staged_erp import live_query

    class_by_code: dict[str, str] = {}
    for row in live_query(_CLASS_BY_CODE_SQL, (normalized,)):
        code = compact_text(row.get("inventory_code")).upper()
        if not code:
            continue
        class_by_code[code] = compact_text(row.get("inventory_class_code")).upper()
    return class_by_code


def score_inventory_candidate(candidate: dict[str, Any], source: dict[str, Any]) -> int:
    code = compact_text(candidate.get("inventory_code"))
    class_code = compact_text(candidate.get("inventory_class_code")).upper()
    score = 0
    if class_code == "RAW MATERIAL":
        score += 100
    if parse_number(candidate.get("remaining_qty"), 0) > 0:
        score += 10
    type_key = _alnum_key(source.get("our_matl_type"))
    code_key = _alnum_key(code)
    if type_key and type_key in code_key:
        score += 50
    od, id_val = parse_inventory_dimensions(code)
    want_od = parse_dimension_number(source.get("matl_od"))
    want_id = parse_dimension_number(source.get("matl_id"))
    if _nums_close(od, want_od):
        score += 30
    if _nums_close(id_val, want_id):
        score += 30
    return score


def pick_inventory(
    candidates: list[dict[str, Any]],
    source: dict[str, Any],
) -> tuple[dict[str, Any] | None, bool]:
    if not candidates:
        return None, False
    ranked = sorted(
        candidates,
        key=lambda item: (
            score_inventory_candidate(item, source),
            parse_number(item.get("remaining_qty"), 0),
            compact_text(item.get("inventory_code")).upper(),
        ),
        reverse=True,
    )
    best = ranked[0]
    if len(ranked) == 1:
        return best, False
    best_score = score_inventory_candidate(best, source)
    second_score = score_inventory_candidate(ranked[1], source)
    return best, best_score == second_score


def _match_status(*, ps_matched: bool, inventory_matched: bool, inventory_ambiguous: bool) -> str:
    if not ps_matched and not inventory_matched:
        return "ps_and_inventory_unmatched"
    if inventory_ambiguous and inventory_matched:
        return "inventory_ambiguous"
    if not ps_matched:
        return "unmatched_ps"
    if not inventory_matched:
        return "unmatched_inventory"
    return "ok"


def match_source_rows(
    source_rows: list[dict[str, Any]],
    aps_by_part: dict[str, list[str]],
    lots_by_ref: dict[str, list[dict[str, Any]]],
    class_by_code: dict[str, str] | None = None,
) -> list[dict[str, Any]]:
    classes = class_by_code or {}
    matched: list[dict[str, Any]] = []
    for source in source_rows:
        part_key = normalize_part_no(source.get("part_no"))
        ps_ids = list(aps_by_part.get(part_key) or [])
        if not ps_ids:
            ps_ids = [""]
        candidates = []
        for item in lots_for_in_house_ref(source.get("in_house_ref"), lots_by_ref):
            row = dict(item)
            code_key = compact_text(row.get("inventory_code")).upper()
            if not row.get("inventory_class_code"):
                row["inventory_class_code"] = classes.get(code_key, "")
            candidates.append(row)
        picked, ambiguous = pick_inventory(candidates, source)
        inventory_code = compact_text((picked or {}).get("inventory_code"))
        remaining = (picked or {}).get("remaining_qty") if picked else ""
        inventory_matched = bool(inventory_code)
        for ps_id in ps_ids:
            row = {field: source.get(field, "") for field in SOURCE_FIELDS}
            row["matched_ps_no"] = compact_text(ps_id)
            row["inventory_code"] = inventory_code
            row["lot_remaining_qty"] = remaining if remaining != "" else ""
            row["match_status"] = _match_status(
                ps_matched=bool(compact_text(ps_id)),
                inventory_matched=inventory_matched,
                inventory_ambiguous=ambiguous,
            )
            matched.append(row)
    return matched


def summarize_rows(matched_rows: list[dict[str, Any]]) -> dict[str, Any]:
    statuses: dict[str, int] = {}
    matched_ps = 0
    unmatched_ps = 0
    unmatched_inventory = 0
    for row in matched_rows:
        status = compact_text(row.get("match_status")) or "ok"
        statuses[status] = statuses.get(status, 0) + 1
        if compact_text(row.get("matched_ps_no")):
            matched_ps += 1
        else:
            unmatched_ps += 1
        if not compact_text(row.get("inventory_code")):
            unmatched_inventory += 1
    return {
        "input_rows": 0,
        "output_rows": len(matched_rows),
        "matched_ps": matched_ps,
        "unmatched_ps": unmatched_ps,
        "unmatched_inventory": unmatched_inventory,
        "statuses": statuses,
    }


def _excel_cell(value: Any) -> Any:
    if value in (None, ""):
        return ""
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        try:
            number = float(value)
        except (TypeError, ValueError):
            return compact_text(value)
        if number != number:
            return ""
        if number.is_integer():
            return int(number)
        return number
    return compact_text(value)


def build_workbook(matched_rows: list[dict[str, Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    ws = workbook.active
    ws.title = "APS match"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="334155")
    wrap = Alignment(horizontal="center", wrap_text=True, vertical="center")
    for index, field in enumerate(EXPORT_FIELDS, start=1):
        cell = ws.cell(row=1, column=index, value=FIELD_LABELS.get(field, field))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap
        ws.column_dimensions[get_column_letter(index)].width = 22 if field != "description" else 36
    for row_idx, record in enumerate(matched_rows, start=2):
        for col_idx, field in enumerate(EXPORT_FIELDS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=_excel_cell(record.get(field)))
    last_col = get_column_letter(len(EXPORT_FIELDS))
    ws.auto_filter.ref = f"A1:{last_col}{max(1, len(matched_rows) + 1)}"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def download_name(when: date | None = None) -> str:
    stamp = (when or planner_today()).isoformat().replace("-", "")
    return f"aps-ps-match-{stamp}.xlsx"


def process_upload(payload: bytes, filename: str = "") -> tuple[bytes, dict[str, Any]]:
    source_rows = parse_workbook(payload, filename)
    aps_by_part = fetch_outstanding_aps()
    refs = [compact_text(row.get("in_house_ref")) for row in source_rows]
    lots_by_ref = fetch_lots_by_reference_nos(refs)
    codes: list[str] = []
    for items in lots_by_ref.values():
        for item in items:
            code = compact_text(item.get("inventory_code"))
            if code:
                codes.append(code)
    class_by_code = fetch_inventory_classes(codes)
    matched = match_source_rows(source_rows, aps_by_part, lots_by_ref, class_by_code)
    summary = summarize_rows(matched)
    summary["input_rows"] = len(source_rows)
    summary["ok"] = True
    summary["download_name"] = download_name()
    return build_workbook(matched), summary
